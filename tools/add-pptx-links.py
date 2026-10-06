# Marp の PPTX は各スライドが画像なのでリンクが効かない。
# measure-links.cjs で測った位置に、透明なクリック領域（ハイパーリンク付き）を重ねる。
import json
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

pptx_path, links_path = sys.argv[1], sys.argv[2]
prs = Presentation(pptx_path)
W, H = prs.slide_width, prs.slide_height

with open(links_path, encoding="utf-8") as f:
    links = json.load(f)

for link in links:
    slide = prs.slides[link["slide"] - 1]
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        int(link["x"] * W), int(link["y"] * H),
        int(link["w"] * W), int(link["h"] * H),
    )
    shape.name = "link"
    shape.fill.background()
    shape.line.fill.background()
    shape.shadow.inherit = False
    shape.click_action.hyperlink.address = link["href"]

prs.save(pptx_path)
print(f"{len(links)} link(s) added to {pptx_path}")
